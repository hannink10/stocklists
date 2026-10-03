#!/usr/bin/env python3
"""Erzeugt aus der aktuellen CSV im Ordner input/ eine Excel-Stockliste
auf Basis von template/stocklist_template.xlsx.

Aufruf:
    python generate_stocklist.py              # CSV aus input/ verwenden
    python generate_stocklist.py --csv X.csv  # bestimmte CSV verwenden
    python generate_stocklist.py --no-archive # CSV nach Erfolg nicht archivieren

Mapping und Regeln stehen in config/mapping.json.
"""

import argparse
import csv
import json
import os
import re
import shutil
import sys
import unicodedata
from collections import Counter, OrderedDict
from copy import copy
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils.units import pixels_to_EMU

BASE_DIR = Path(__file__).resolve().parent
CONFIG_PATH = BASE_DIR / "config" / "mapping.json"


class StocklistError(Exception):
    """Kritischer Fehler: es wird keine Kundendatei erzeugt."""


# ---------------------------------------------------------------------------
# Eingabe
# ---------------------------------------------------------------------------

def load_config():
    with open(CONFIG_PATH, encoding="utf-8") as f:
        return json.load(f)


def find_input_csv(cfg, explicit):
    if explicit:
        path = Path(explicit)
        if not path.is_file():
            raise StocklistError(f"CSV-Datei nicht gefunden: {path}")
        return path

    input_dir = BASE_DIR / cfg["paths"]["input_dir"]
    csvs = sorted(p for p in input_dir.glob("*") if p.suffix.lower() == ".csv")
    if not csvs:
        raise StocklistError(f"Keine CSV-Datei im Ordner '{input_dir}' gefunden.")
    if len(csvs) > 1:
        names = "\n".join(f"  - {p.name}" for p in csvs)
        raise StocklistError(
            "Im Input-Ordner liegen mehrere CSV-Dateien:\n" + names + "\n"
            "Bitte nur die aktuelle CSV im Ordner lassen (alte nach input/archiv verschieben) "
            "oder die gewünschte Datei angeben: python generate_stocklist.py --csv input/DATEI.csv"
        )
    return csvs[0]


def detect_delimiter(path, enc, configured):
    """Trennzeichen: fest aus mapping.json oder 'auto' (das in den ersten Zeilen häufigste von ; , Tab |)."""
    if configured != "auto":
        return configured
    with open(path, encoding=enc, errors="strict") as f:
        head = "".join(f.readline() for _ in range(5))
    counts = {d: head.count(d) for d in (";", ",", "\t", "|")}
    return max(counts, key=counts.get)


def read_csv(path, cfg):
    """Liest die CSV tolerant: unterschiedlich lange Zeilen (z.B. Titelzeilen über der Kopfzeile) sind erlaubt.
    Zeilennummern bleiben wie in der Datei (Leerzeilen werden nicht entfernt)."""
    c = cfg["csv"]
    last_error = None
    for enc in c["encodings"]:
        try:
            sep = detect_delimiter(path, enc, c["delimiter"])
            with open(path, encoding=enc, newline="") as f:
                rows = list(csv.reader(f, delimiter=sep))
        except UnicodeDecodeError as e:
            last_error = e
            continue
        except csv.Error as e:
            raise StocklistError(f"CSV konnte nicht gelesen werden (Format/Trennzeichen?): {e}")
        if not any(any(v.strip() for v in r) for r in rows):
            raise StocklistError(f"Die CSV-Datei ist leer: {path}")
        width = max(len(r) for r in rows)
        # z.B. 'E' + Akzent als ein Zeichen 'É' speichern
        rows = [[unicodedata.normalize("NFC", v) for v in r] + [""] * (width - len(r)) for r in rows]
        return pd.DataFrame(rows, dtype=str), enc
    raise StocklistError(f"CSV-Encoding nicht erkannt ({', '.join(c['encodings'])}): {last_error}")


def parse_number(text, cfg):
    """Zahl aus der CSV -> Decimal. Versteht '1.234,50', '1,234.50', '62,94', '62.94', '€ 62,94', '3,00-'.
    Bei nur einem Trennzeichen entscheidet mapping.json (csv.decimal), außer es folgen genau 3 Ziffern
    nicht im Sinne von Nachkommastellen (z.B. '1.000' bei Komma als Dezimalzeichen = Tausend)."""
    s = text.strip().replace("\u2212", "-").replace("\u00a0", "").replace(" ", "")
    s = re.sub(r"[€$£]|EUR|USD|CHF", "", s, flags=re.I)
    if s == "":
        return None
    dec = cfg["csv"].get("decimal", ",")
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        if sep != dec and not re.fullmatch(r"-?\d{1,3}(" + re.escape(sep) + r"\d{3})+-?", s):
            dec = sep  # z.B. '62.94' obwohl Komma erwartet: eindeutig Dezimalpunkt
    thousands = "." if dec == "," else ","
    s = s.replace(thousands, "").replace(dec, ".")
    if s.endswith("-"):  # nachgestelltes Minus, z.B. '3,00-'
        s = "-" + s[:-1].strip()
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(text)


def parse_discount(text, cfg):
    """Discount aus der CSV -> Anteil (0.30 für 30 %). Versteht '30', '30%', '30,0 %', '0,3'. Leer -> None.
    Werte unter 1 ohne %-Zeichen gelten als Anteil (0,3 = 30 %), sonst als Prozent."""
    s = text.strip()
    if s == "":
        return None
    value = parse_number(s.replace("%", ""), cfg)
    if value is None:
        return None
    frac = value / 100 if ("%" in s or value >= 1) else value
    if frac < 0 or frac >= 1:
        raise ValueError(text)
    return frac


def cents(value):
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# ---------------------------------------------------------------------------
# Prüfung & Aufbereitung
# ---------------------------------------------------------------------------

def brand_config(cfg, brand):
    """Einstellungen für eine Marke aus mapping.json → brands (Name wie in config/shopify.env)."""
    for name, bcfg in cfg.get("brands", {}).items():
        if not name.startswith("_") and isinstance(bcfg, dict) and normalize_name(name) == normalize_name(brand or ""):
            return bcfg
    return {}


def apply_brand_groups(cfg, brand):
    """Produktgruppen je Marke: zusätzliche Stichwörter (nach den allgemeinen geprüft) und eigene Reihenfolge."""
    bg = brand_config(cfg, brand).get("product_groups", {})
    gcfg = cfg.setdefault("product_groups", {})
    gcfg["_brand"] = brand
    gcfg["learned"] = load_learned_groups(brand)
    if bg.get("rules"):
        gcfg["rules"] = list(gcfg.get("rules", [])) + list(bg["rules"])
    if bg.get("order"):
        gcfg["order"] = list(bg["order"])
    if bg.get("sku_prefixes"):
        gcfg["sku_prefixes"] = list(bg["sku_prefixes"])


def split_sku(sku, rule, known_sizes):
    """SKU -> (Artikelnummer, Größe, umgerechnet). Größe None = Artikel ohne Größe (One-Size).
    rule 'separator': Größe nach dem letzten Trennzeichen ('1032212-XS' -> '1032212', 'XS').
    rule 'patterns':  Liste regulärer Ausdrücke mit den Gruppen 'base' und 'size' – als Text oder als
                      {"pattern": ..., "size_map": {...}} mit Umrechnung (z.B. US 'W85' -> EU '40');
                      das erste Muster, das passt und eine bekannte Größe liefert, gewinnt."""
    if rule.get("mode") == "patterns":
        for entry in rule["patterns"]:
            pattern, size_map = (entry, None) if isinstance(entry, str) else (entry["pattern"], entry.get("size_map"))
            m = re.fullmatch(pattern, sku)
            if not m:
                continue
            size = m.group("size")
            if size_map is not None:
                size = size_map.get(size)
            if size in known_sizes:
                return m.group("base"), size, size_map is not None
        return sku, None, False
    sep = rule.get("separator", "-")
    if sep in sku:
        base, size = sku.rsplit(sep, 1)
        return base, size, False
    return sku, None, False


def norm_header(text):
    """Spaltenname zum Vergleichen: Großbuchstaben, ohne Akzente/Umlaute-Punkte, Satzzeichen und doppelte Leerzeichen."""
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


ROLE_LABELS = {"sku": "SKU", "name": "Name", "uvp": "UVP/RRP", "unit_price": "Händlerpreis", "stock": "Bestand",
               "discount": "Discount"}


def column_letter_csv(i):
    return get_column_letter(i + 1)


def detect_columns(raw, cfg, brand, warnings):
    """Findet Kopfzeile und Spalten über ihre Namen (mit Synonymen aus mapping.json, je Marke überschreibbar).
    Liefert (Zeilenindex der Kopfzeile, {Feld: Spaltenindex}, Beschreibung für die Ausgabe)."""
    columns = {k: v for k, v in cfg["columns"].items() if not k.startswith("_")}
    brand_cols = brand_config(cfg, brand).get("columns", {})

    def aliases(role):
        spec = dict(columns[role], **brand_cols.get(role, {}))
        if role in brand_cols and "header" in brand_cols[role]:
            names = [brand_cols[role]["header"]]  # Marke gibt den Namen fest vor
        else:
            names = [spec["header"]] + spec.get("aliases", [])
        return [norm_header(n) for n in names], spec

    # Kopfzeile = Zeile (unter den ersten 15) mit den meisten erkannten Feldern
    best_idx, best_hits = 0, -1
    for idx in range(min(15, len(raw))):
        cells = {norm_header(v) for v in raw.iloc[idx].tolist()}
        hits = sum(any(a in cells for a in aliases(role)[0]) for role in columns)
        if hits > best_hits:
            best_idx, best_hits = idx, hits
    header = [str(v).strip() for v in raw.iloc[best_idx].tolist()]
    normed = [norm_header(h) for h in header]

    positions, used, info, errors = {}, set(), [], []
    for role in columns:
        names, spec = aliases(role)
        hits = []
        for n in names:  # Synonyme in Reihenfolge: das erste, das vorkommt, gewinnt
            hits = [i for i, h in enumerate(normed) if h == n and i not in used]
            if hits:
                break
        if not hits:
            if spec.get("required", True):
                errors.append(f"Keine Spalte für '{ROLE_LABELS.get(role, role)}' gefunden "
                              f"(gesucht: {', '.join(spec_names(columns[role], brand_cols.get(role)))}).")
            continue
        pick, note = hits[0], ""
        if len(hits) > 1:
            values = [tuple(raw.iloc[best_idx + 1:, i].tolist()) for i in hits]
            occ = spec.get("occurrence", 1)
            if len(set(values)) == 1:
                note = f" (von {len(hits)} gleichen Spalten)"
            elif 1 <= occ <= len(hits):
                pick, note = hits[occ - 1], f" ({occ}. von {len(hits)} Spalten '{header[hits[0]]}')"
            else:
                pick = hits[-1]
                note = f" (letzte von {len(hits)} Spalten '{header[hits[0]]}')"
                warnings.append(f"Spalte '{header[hits[0]]}' kommt {len(hits)}x mit unterschiedlichen Werten vor – "
                                f"verwendet die letzte ({column_letter_csv(pick)}). Bitte prüfen.")
        positions[role] = pick
        used.add(pick)
        info.append(f"{ROLE_LABELS.get(role, role)} = '{header[pick]}' ({column_letter_csv(pick)}){note}")

    if errors:
        found = ", ".join(f"'{h}'" for h in header if h)
        raise StocklistError(
            "Spaltenprüfung fehlgeschlagen:\n  - " + "\n  - ".join(errors)
            + f"\n\nSpalten in der CSV (Zeile {best_idx + 1}): {found}\n"
            "Lösung: in config/mapping.json den Spaltennamen unter columns.<feld>.aliases ergänzen "
            "oder für diese Marke unter brands festlegen (siehe README)."
        )
    return best_idx, positions, info


def spec_names(spec, brand_spec=None):
    if brand_spec and "header" in brand_spec:
        return [brand_spec["header"]]
    return [spec["header"]] + spec.get("aliases", [])


def build_articles(raw, positions, cfg, template_sizes, warnings, header_idx=0, brand=None, markup=None):
    """markup: EK-Preis = UVP ÷ markup (nur wenn die CSV keine EK-Spalte hat)."""
    rule = brand_config(cfg, brand).get("sku_rule") or {"mode": "separator", "separator": cfg["sizes"]["sku_separator"]}
    no_size = cfg["sizes"]["no_size_column"]
    extra = [e["size"] for e in cfg["sizes"]["extra_sizes"]]
    known_sizes = set(template_sizes) | set(extra)

    errors = []
    negative = []
    no_size_skus = []       # SKU ohne erkennbare Größe (Suffix-Regel) -> One-Size
    merged = []             # umgerechnete Größen, die doppelt vorkommen (Bestand addiert)
    name_without_size = []  # Name endet nicht auf die Größe -> Name unverändert übernommen
    seen_skus = {}
    articles = OrderedDict()

    for idx in range(header_idx + 1, len(raw)):
        row = raw.iloc[idx]
        line = idx + 1  # Zeilennummer in der CSV-Datei
        val = {k: str(row.iloc[p]).strip() for k, p in positions.items()}

        if all(v == "" for v in val.values()):
            continue
        required = [k for k in ("sku", "name", "uvp", "unit_price", "stock") if k in positions]
        for k in required:
            if val[k] == "":
                errors.append(f"Zeile {line}: Pflichtfeld '{ROLE_LABELS.get(k, k)}' ist leer.")
        if any(val[k] == "" for k in required):
            continue

        sku = val["sku"]
        if sku in seen_skus:
            errors.append(f"Zeile {line}: SKU '{sku}' doppelt (bereits in Zeile {seen_skus[sku]}).")
            continue
        seen_skus[sku] = line

        try:
            uvp = cents(parse_number(val["uvp"], cfg))  # Preise immer auf 2 Nachkommastellen
            price = cents(parse_number(val["unit_price"], cfg)) if "unit_price" in positions else cents(uvp / markup)
            stock = parse_number(val["stock"], cfg)
            discount = parse_discount(val["discount"], cfg) if "discount" in positions else None
        except ValueError as e:
            errors.append(f"Zeile {line} (SKU {sku}): ungültiger Wert '{e}'.")
            continue
        if stock != stock.to_integral_value():
            errors.append(f"Zeile {line} (SKU {sku}): ungültiger Bestand '{val['stock']}'.")
            continue
        if stock < 0:
            # negativer Bestand (z.B. überverkauft) zählt als 0
            negative.append(f"{sku} ({int(stock)})")
            stock = Decimal(0)
        if price <= 0 or uvp <= 0:
            errors.append(f"Zeile {line} (SKU {sku}): Preis <= 0.")
            continue
        if price > uvp:
            warnings.append(f"Zeile {line} (SKU {sku}): UNIT PRICE {price} ist höher als UVP {uvp}.")

        # Artikelnummer und Größe trennen (Regel je Marke, siehe mapping.json → brands → sku_rule)
        base, size, converted = split_sku(sku, rule, known_sizes)
        if size is None:
            size, name = no_size, val["name"]
            if rule.get("mode") == "patterns":
                no_size_skus.append(sku)
        else:
            if size not in known_sizes:
                errors.append(
                    f"Zeile {line}: unbekannte Größe '{size}' in SKU '{sku}'. "
                    "Bitte in mapping.json unter sizes.extra_sizes ergänzen."
                )
                continue
            name = re.sub(r"\s*[-/]?\s*" + re.escape(size) + r"$", "", val["name"]).strip()
            if name == val["name"] and converted:
                # umgerechnete Größe: im Namen steht meist die Ursprungsgröße, z.B. 'FORGET ME NOT MEN - 7,5'
                name = re.sub(r"\s*[-/]\s*\d{1,2}(?:[,.]\d)?$", "", val["name"]).strip()
            if converted:
                # Größensystem-Zusatz entfernen: 'SEED.ONE - MEMO-Z - EU Women' -> 'SEED.ONE - MEMO-Z'
                stripped = re.sub(r"\s*[-/]\s*EU(\s+(WOMEN|WMNS|MEN))?$", "", name, flags=re.I).strip()
                if stripped != name:
                    name = stripped
            if name == val["name"]:
                name_without_size.append(sku)

        art = articles.get(base)
        if art is None:
            art = articles[base] = {
                "sku": base, "name": name, "uvp": uvp, "unit_price": price,
                "stock": OrderedDict(), "skus": [], "first_line": line,
                "image": None, "sold": 0, "prices": {"unit_price": [], "uvp": [], "discount": []},
            }
        art.setdefault("names", []).append(name)
        art["prices"]["unit_price"].append((price, size))
        art["prices"]["uvp"].append((uvp, size))
        art["prices"]["discount"].append((discount or Decimal(0), size))
        if size in art["stock"] and converted:
            # z.B. US Damen 8,5 und US Herren 7 = beide EU 40: Bestand zusammenzählen
            merged.append(f"{base} {size}")
            art["stock"][size] += int(stock)
            art["skus"].append(sku)
            continue
        if size in art["stock"]:
            errors.append(f"Zeile {line}: Artikel {base} hat Größe '{size}' mehrfach.")
        art["stock"][size] = int(stock)
        art["skus"].append(sku)

    resolve_name_conflicts(list(articles.values()), warnings)
    resolve_price_conflicts(list(articles.values()), cfg, warnings)
    for art in articles.values():
        art["discount_price"] = cents(art["unit_price"] * (1 - art["discount"]))
    if no_size_skus:
        warnings.append(f"{len(no_size_skus)} SKU(s) ohne erkennbare Größe – als One-Size übernommen "
                        "(falls falsch: Muster in mapping.json → brands → sku_rule ergänzen): "
                        + ", ".join(no_size_skus[:15]) + (" …" if len(no_size_skus) > 15 else ""))
    if merged:
        warnings.append(f"{len(merged)} umgerechnete Größe(n) kamen doppelt vor (z.B. US Damen + Herren = gleiche EU-Größe) "
                        "– Bestand zusammengezählt: " + ", ".join(merged[:10]) + (" …" if len(merged) > 10 else ""))
    if name_without_size:
        warnings.append(f"{len(name_without_size)} Zeile(n): Name endet nicht auf die Größe (Name unverändert übernommen), z.B. "
                        + ", ".join(name_without_size[:5]))
    if negative:
        warnings.append(f"{len(negative)} Größe(n) mit negativem Bestand auf 0 gesetzt: " + ", ".join(negative))
    if errors:
        shown = errors[:30]
        more = f"\n  ... und {len(errors) - 30} weitere" if len(errors) > 30 else ""
        raise StocklistError("Datenprüfung fehlgeschlagen:\n  - " + "\n  - ".join(shown) + more)
    if not articles:
        raise StocklistError("Die CSV enthält keine Artikel.")

    result = list(articles.values())
    if cfg["rules"]["skip_articles_without_stock"]:
        skipped = [a for a in result if sum(a["stock"].values()) == 0]
        result = [a for a in result if sum(a["stock"].values()) > 0]
        if skipped:
            warnings.append(
                f"{len(skipped)} Artikel ohne Bestand weggelassen: "
                + ", ".join(f"{a['sku']} ({a['name']})" for a in skipped)
            )
    return result


def model_name(name):
    """Produktname ohne Farbe: 'QUOTE T-SHIRT - CREAM' -> 'QUOTE T-SHIRT'."""
    return name.rsplit(" - ", 1)[0].strip() if " - " in name else name


def common_name(names):
    """Gemeinsamer Namensanfang an einer Wortgrenze, ohne angehängte Trenner und 'EU'
    ('SEED.ONE - MEMO-Z - EU Women' + 'SEED.ONE - MEMO-Z - EU' -> 'SEED.ONE - MEMO-Z')."""
    prefix = os.path.commonprefix(names)
    if any(len(n) > len(prefix) and n[len(prefix)] not in " -/" for n in names):
        prefix = prefix[:prefix.rfind(" ")] if " " in prefix else ""
    prefix = re.sub(r"(\s*[-/]\s*|\s+)(EU)?\s*$", "", prefix).strip(" -/")
    prefix = re.sub(r"\s*[-/]\s*EU$", "", prefix).strip(" -/")
    return prefix


def resolve_name_conflicts(articles, warnings):
    """Haben die Größen eines Artikels unterschiedliche Namen (z.B. '… - EU Women' und '… - EU'),
    gilt der gemeinsame Namensanfang, sonst der häufigste Name. Kein Abbruch, aber Hinweis."""
    changed = []
    for art in articles:
        raw = art.pop("names", [art["name"]])
        names = list(OrderedDict.fromkeys(raw))
        if len(names) == 1:
            art["name"] = names[0]
            continue
        chosen = common_name(names)
        if len(chosen) < 3:
            chosen = Counter(raw).most_common(1)[0][0]
        art["name"] = chosen
        changed.append(f"{art['sku']}: " + " / ".join(f"'{n}'" for n in names) + f" → '{chosen}'")
    if changed:
        warnings.append(f"{len(changed)} Artikel mit unterschiedlichen Namen je Größe – gemeinsamer Name verwendet: "
                        + "; ".join(changed[:5]) + (" …" if len(changed) > 5 else ""))


def resolve_price_conflicts(articles, cfg, warnings):
    """Haben die Größen eines Artikels unterschiedliche Preise, gilt der häufigste Preis.
    Bei Gleichstand entscheidet der Preis, den dasselbe Modell in anderen Farben hat;
    ist es dann immer noch offen, der höhere Preis. Jede Abweichung wird als Hinweis gemeldet."""
    labels = {"unit_price": "EK-Preis", "uvp": "UVP", "discount": "Discount"}
    by_model = {}
    for art in articles:
        by_model.setdefault(model_name(art["name"]), []).append(art)

    for art in articles:
        for key, entries in art["prices"].items():
            def fmt(v, key=key):
                if key == "discount":
                    return f"{v * 100:.0f} %"
                return f"{v:.2f}".replace(".", ",") + " €"

            counts = Counter(v for v, _ in entries)
            if len(counts) == 1:
                art[key] = entries[0][0]
                continue
            top = max(counts.values())
            candidates = [v for v, c in counts.items() if c == top]
            reason = "häufigster Preis"
            if len(candidates) > 1:
                siblings = Counter(
                    v for other in by_model[model_name(art["name"])] if other is not art
                    for v, _ in other["prices"][key] if v in candidates
                )
                if siblings:
                    best = max(siblings.values())
                    in_siblings = [v for v in candidates if siblings[v] == best]
                    if len(in_siblings) == 1:
                        candidates, reason = in_siblings, "Preis des Modells in anderen Farben"
            if len(candidates) > 1:
                reason = "kein eindeutiger Preis, höherer Preis genommen"
            chosen = max(candidates)
            art[key] = chosen
            detail = "; ".join(
                f"{fmt(v)}: " + ", ".join(sz for pv, sz in entries if pv == v)
                for v in sorted(counts)
            )
            warnings.append(
                f"Artikel {art['sku']} ({art['name']}): unterschiedliche {labels[key]} je Größe "
                f"({detail}) – verwendet {fmt(chosen)} ({reason}). Bitte in JTL prüfen."
            )


OVERRIDES_FILE = BASE_DIR / "config" / "category_overrides.json"


def load_learned_groups(brand):
    """Im Fenster zugeordnete Kategorien (config/category_overrides.json): {Marke: {Modellname: Gruppe}}."""
    try:
        data = json.loads(OVERRIDES_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    for name, groups in data.items():
        if normalize_name(name) == normalize_name(brand or "") and isinstance(groups, dict):
            return {normalize_name(k): v for k, v in groups.items()}
    return {}


def save_learned_groups(brand, new):
    """Neue Zuordnungen {Modellname: Gruppe} dauerhaft speichern (bestehende bleiben erhalten)."""
    try:
        data = json.loads(OVERRIDES_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        data = {}
    key = next((k for k in data if normalize_name(k) == normalize_name(brand or "")), brand or "Alle")
    data.setdefault(key, {}).update(new)
    OVERRIDES_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def product_group(name, gcfg, sku=None):
    """Produktgruppe: zuerst eigene Zuordnung aus dem Fenster (category_overrides.json), dann SKU-Anfang
    (sku_prefixes, z.B. '40' = Schuhe), dann Stichwort im Artikelnamen (ohne Farbe); erste passende Regel gewinnt."""
    learned = gcfg.get("learned", {}).get(normalize_name(model_name(name)))
    if learned:
        return learned
    for prefix, group in gcfg.get("sku_prefixes", []):
        if sku and str(sku).upper().startswith(str(prefix).upper()):
            return group
    text = " " + normalize_name(model_name(name)) + " "
    for keyword, group in gcfg["rules"]:
        if " " + normalize_name(keyword) + " " in text:
            return group
    return gcfg["other_group"]


def sort_articles(articles, cfg, warnings):
    """Sortierung (stabil: bei Gleichstand bleibt die CSV-Reihenfolge).
    Mit Produktgruppen: erst die meistverkauften Artikel (Abschnitt Bestseller), dann je Gruppe.
    Setzt art["section"] = Überschrift des Abschnitts (None ohne Gruppen)."""
    mode = cfg["rules"]["sort"]
    if mode == "stock_desc":
        articles.sort(key=lambda a: -sum(a["stock"].values()))
    elif mode == "bestseller":
        articles.sort(key=lambda a: (-a["sold"], -sum(a["stock"].values())))
    elif mode != "csv":
        raise StocklistError(f"Unbekannte Sortierung '{mode}' in mapping.json (stock_desc, bestseller, csv).")

    gcfg = cfg.get("product_groups", {})
    if not gcfg.get("enabled"):
        for art in articles:
            art["section"] = None
        return

    top = []
    if mode == "bestseller" and gcfg.get("bestseller_top"):
        top = [a for a in articles if a["sold"] > 0][: gcfg["bestseller_top"]]
    for art in top:
        art["section"] = gcfg["bestseller_title"]
    top_ids = {id(a) for a in top}

    order = list(gcfg["order"])
    rest = [a for a in articles if id(a) not in top_ids]
    for art in rest:
        art["section"] = product_group(art["name"], gcfg, art["sku"])
        if art["section"] not in order:
            order.append(art["section"])
    rest.sort(key=lambda a: order.index(a["section"]))  # stabil: innerhalb der Gruppe bleibt die Sortierung oben

    other = [a for a in rest if a["section"] == gcfg["other_group"]]
    if other and gcfg.get("ask_unknown", True) and sys.stdin.isatty():
        ask_unknown_groups(other, gcfg)
        for art in other:
            art["section"] = product_group(art["name"], gcfg, art["sku"])
            if art["section"] not in order:
                order.append(art["section"])
        rest.sort(key=lambda a: order.index(a["section"]))
        other = [a for a in rest if a["section"] == gcfg["other_group"]]
    if other:
        warnings.append(
            f"{len(other)} Artikel keiner Produktgruppe zugeordnet ('{gcfg['other_group']}'): "
            + ", ".join(f"{a['sku']} ({a['name']})" for a in other)
            + " – beim nächsten Start im Fenster zuordnen oder Stichwort in mapping.json ergänzen."
        )
    articles[:] = top + rest


def ask_unknown_groups(other, gcfg):
    """Fragt für Artikel ohne Kategorie nach der Gruppe und speichert die Antwort dauerhaft (je Modellname,
    gilt damit auch für andere Farben desselben Modells)."""
    groups = [g for g in gcfg["order"] if g != gcfg["other_group"]]
    models = OrderedDict()
    for art in other:
        models.setdefault(model_name(art["name"]), art)
    print(f"{len(models)} Artikel ohne Kategorie – bitte zuordnen (wird für die Zukunft gespeichert):")
    print("    " + "  ".join(f"{i} {g}" for i, g in enumerate(groups, 1)) + f"  0 = {gcfg['other_group']} lassen")
    new = {}
    for model, art in models.items():
        while True:
            answer = input(f"  {art['sku']}  {model}: ").strip()
            if answer.isdigit() and 0 <= int(answer) <= len(groups):
                break
            print(f"    Bitte eine Zahl von 0 bis {len(groups)} eingeben.")
        if int(answer) > 0:
            new[model] = groups[int(answer) - 1]
    print()
    if new:
        save_learned_groups(gcfg.get("_brand"), new)
        gcfg.setdefault("learned", {}).update({normalize_name(k): v for k, v in new.items()})
        print(f"Gespeichert in config/{OVERRIDES_FILE.name}: " + ", ".join(f"{k} → {v}" for k, v in new.items()) + "\n")


def load_shops(cfg):
    try:
        import shopify_client as sc
    except ImportError:
        return []
    return sc.load_shops(BASE_DIR / "config" / "shopify.env", cfg["excel"]["default_brand"])


def select_shop(shops, wanted):
    """Shop/Marke für diesen Lauf: per --shop, bei mehreren Shops sonst per Auswahl im Fenster."""
    if wanted:
        key = normalize_name(wanted)
        for s in shops:
            if key in (normalize_name(s["name"]), normalize_name(s["shop"]), normalize_name(s["shop"].split(".")[0])):
                return s
        raise StocklistError(f"Shop '{wanted}' nicht in config/shopify.env. Vorhanden: "
                             + ", ".join(s["name"] for s in shops))
    if len(shops) <= 1:
        return shops[0] if shops else None
    if not sys.stdin.isatty():
        raise StocklistError("Mehrere Shops in config/shopify.env – bitte mit --shop <Marke> wählen: "
                             + ", ".join(s["name"] for s in shops))
    print("Für welche Marke soll die Stockliste erstellt werden?")
    for i, s in enumerate(shops, 1):
        print(f"  {i}  {s['name']}")
    while True:
        answer = input("Nummer eingeben und Enter drücken: ").strip()
        if answer.isdigit() and 1 <= int(answer) <= len(shops):
            print()
            return shops[int(answer) - 1]
        print(f"Bitte eine Zahl von 1 bis {len(shops)} eingeben.")


SORT_CHOICES = [
    ("bestseller", "Bestseller & Kategorien"),
    ("categories", "Nur Kategorien (je Kategorie: meistverkauft zuerst, dann höchster Bestand)"),
    ("stock", "Nach Bestand (höchster Bestand zuerst, ohne Kategorien)"),
]


def select_sort(wanted, cfg):
    """Sortierung für diesen Lauf: per --sort, sonst Auswahl im Fenster, ohne Fenster Standard aus mapping.json."""
    keys = [k for k, _ in SORT_CHOICES]
    if wanted:
        choice = wanted
    elif sys.stdin.isatty():
        print("Wie soll sortiert werden?")
        for i, (_, label) in enumerate(SORT_CHOICES, 1):
            print(f"  {i}  {label}")
        while True:
            answer = input("Nummer eingeben und Enter drücken: ").strip()
            if answer.isdigit() and 1 <= int(answer) <= len(keys):
                choice = keys[int(answer) - 1]
                print()
                break
            print(f"Bitte eine Zahl von 1 bis {len(keys)} eingeben.")
    else:
        choice = "bestseller" if cfg["rules"]["sort"] == "bestseller" else "stock"
    gcfg = cfg.setdefault("product_groups", {})
    if choice == "stock":
        cfg["rules"]["sort"] = "stock_desc"
        gcfg["enabled"] = False
    else:
        # Verkaufszahlen + Bestand sortieren auch innerhalb der Kategorien; 'categories' ohne Bestseller-Block oben
        cfg["rules"]["sort"] = "bestseller"
        gcfg["enabled"] = True
        if choice == "categories":
            gcfg["bestseller_top"] = 0
    return dict(SORT_CHOICES)[choice]


def ask_markup(wanted):
    """Markup für den EK-Preis (EK = UVP ÷ Markup): per --markup, sonst Abfrage im Fenster."""
    def parse(text):
        try:
            value = Decimal(text.strip().replace(",", "."))
        except InvalidOperation:
            return None
        return value if value >= 1 else None

    if wanted:
        value = parse(wanted)
        if value is None:
            raise StocklistError(f"Ungültiges Markup '{wanted}' (z.B. 2,5).")
        return value
    if not sys.stdin.isatty():
        raise StocklistError("Die CSV hat keinen EK-Preis. Bitte Markup angeben, z.B. --markup 2,5")
    print("Die CSV hat keinen EK-Preis – er wird aus der UVP berechnet (EK = UVP ÷ Markup).")
    while True:
        value = parse(input("Markup eingeben (z.B. 2,5) und Enter drücken: "))
        if value is not None:
            print()
            return value
        print("Bitte eine Zahl ab 1 eingeben, z.B. 2,5.")


def normalize_name(name):
    """Name für den Abgleich mit Shopify: ohne Akzente, Satzzeichen und Groß-/Kleinschreibung."""
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]+", " ", name).strip()


def pick_image(candidates, name):
    """Bild-URL aus mehreren Shopify-Treffern: gleicher Name vor anderem Namen, aktiv vor archiviert."""
    if not candidates:
        return None
    key = normalize_name(name)
    best = min(candidates, key=lambda v: (normalize_name(v["title"]) != key, v["status"] != "ACTIVE"))
    return best["url"]


def enrich_from_shopify(articles, cfg, warnings, shop):
    """Bilder und Verkaufszahlen aus Shopify. Fehler hier sind nie kritisch."""
    scfg = cfg.get("shopify", {})
    if not scfg.get("enabled") or shop is None or not (shop["client_id"] and shop["client_secret"]):
        if scfg.get("enabled"):
            warnings.append("Keine Shopify-Zugangsdaten (config/shopify.env) – Liste ohne Bilder, sortiert nach Bestand.")
        if cfg["rules"]["sort"] == "bestseller":
            cfg["rules"]["sort"] = "stock_desc"
        return
    try:
        import shopify_client as sc
    except ImportError as e:
        warnings.append(f"Shopify-Modul nicht ladbar ({e}) – Liste ohne Bilder/Bestseller.")
        cfg["rules"]["sort"] = "stock_desc" if cfg["rules"]["sort"] == "bestseller" else cfg["rules"]["sort"]
        return
    client = sc.ShopifyClient(shop, scfg["api_version"])

    if scfg.get("images"):
        try:
            variants = client.image_variants()
            by_sku, by_name = {}, {}
            for v in variants:
                by_sku.setdefault(v["sku"], []).append(v)
                by_name.setdefault(normalize_name(v["title"]), []).append(v)
            cache = BASE_DIR / scfg["image_cache_dir"]
            missing, via_name = [], []
            for art in articles:
                # 1. SKU mit Größe, 2. SKU ohne Größe (so sind viele Artikel in Shopify angelegt)
                cands = [v for s in art["skus"] + [art["sku"]] for v in by_sku.get(s, [])]
                if not cands and scfg.get("match_by_name", True):
                    # 3. Produktname
                    cands = by_name.get(normalize_name(art["name"]), [])
                    if cands:
                        via_name.append(f"{art['sku']} ({art['name']})")
                url = pick_image(cands, art["name"])
                if url is None:
                    missing.append(art["sku"])
                    continue
                try:
                    art["image"] = sc.thumbnail(url, cache, art["sku"], scfg["image_max_px"])
                except (sc.ShopifyError, OSError) as e:
                    missing.append(art["sku"])
                    warnings.append(f"Bild für {art['sku']} nicht ladbar: {e}")
            if via_name:
                warnings.append(f"{len(via_name)} Artikel-Bild(er) über den Produktnamen zugeordnet "
                                "(SKU nicht in Shopify): " + ", ".join(via_name))
            if missing:
                more = f" … (+{len(missing) - 10} weitere)" if len(missing) > 10 else ""
                warnings.append(f"{len(missing)} Artikel ohne Shopify-Bild: " + ", ".join(missing[:10]) + more)
        except sc.ShopifyError as e:
            warnings.append(f"Shopify-Bilder nicht abrufbar: {e}")

    if cfg["rules"]["sort"] == "bestseller":
        try:
            sold = client.units_sold_by_sku(scfg["bestseller_days"])
            for art in articles:
                # auch Verkäufe unter der SKU ohne Größe zählen
                art["sold"] = sum(sold.get(s, 0) for s in set(art["skus"]) | {art["sku"]})
        except sc.ShopifyError as e:
            warnings.append(f"Shopify-Verkaufszahlen nicht abrufbar ({e}) – sortiere nach Bestand.")
            cfg["rules"]["sort"] = "stock_desc"


def add_image(ws, path, row, col, box_w, box_h):
    """Bild zentriert in die (verbundene) Zelle setzen."""
    img = XLImage(str(path))
    scale = min(box_w / img.width, box_h / img.height, 1)
    w, h = int(img.width * scale), int(img.height * scale)
    img.width, img.height = w, h
    marker = AnchorMarker(col=col - 1, colOff=pixels_to_EMU((box_w - w) // 2 + 4),
                          row=row - 1, rowOff=pixels_to_EMU((box_h - h) // 2 + 4))
    img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
    ws.add_image(img)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def read_template_layout(ws, cfg):
    """Liest Spaltenrollen, Stile, Breiten und Höhen aus dem Template."""
    xc = cfg["excel"]
    hr, r1 = xc["header_row"], xc["first_data_row"]
    r2 = r1 + 1
    headers = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(hr, col).value
        if v is not None:
            headers[str(v).strip()] = col

    col_of = {}
    for role, title in xc["column_headers"].items():
        if title not in headers:
            raise StocklistError(f"Template: Überschrift '{title}' nicht in Zeile {hr} gefunden.")
        col_of[role] = headers[title]
    first_size, last_size = col_of["uvp"] + 1, col_of["quantity"] - 1
    sizes = [str(ws.cell(hr, c).value).strip() for c in range(first_size, last_size + 1)]
    if cfg["sizes"]["no_size_column"] not in sizes:
        raise StocklistError(f"Template: Größenspalte '{cfg['sizes']['no_size_column']}' fehlt.")
    fixed = sorted(col_of[r] for r in ("image", "sku", "name", "unit_price", "uvp"))
    if fixed != list(range(1, len(fixed) + 1)) or col_of["total"] != col_of["quantity"] + 1:
        raise StocklistError("Template: unerwartete Spaltenanordnung.")

    def width(col):
        letter = get_column_letter(col)
        for dim in ws.column_dimensions.values():
            if dim.min and dim.max and dim.min <= col <= dim.max:
                return dim.width
        return ws.column_dimensions[letter].width

    def cell_info(row, col):
        c = ws.cell(row, col)
        return {"style": c._style, "value": c.value}

    roles = {}
    for role in ("image", "sku", "name", "unit_price", "uvp", "quantity", "total"):
        col = col_of[role]
        roles[role] = {
            "header": cell_info(hr, col), "top": cell_info(r1, col),
            "bottom": cell_info(r2, col), "width": width(col),
        }
    roles["size"] = {
        "header": cell_info(hr, first_size), "top": cell_info(r1, first_size),
        "bottom": cell_info(r2, first_size), "width": width(first_size),
    }
    return {
        "sizes": sizes,
        "fixed_order": [r for r in ("image", "sku", "name", "unit_price", "uvp")],
        "fixed_cols": {r: col_of[r] for r in ("image", "sku", "name", "unit_price", "uvp")},
        "roles": roles,
        "row_heights": {
            "title": ws.row_dimensions[1].height,
            "header": ws.row_dimensions[hr].height,
            "top": ws.row_dimensions[r1].height,
            "bottom": ws.row_dimensions[r2].height,
        },
        "title": ws.cell(1, 1).value,
    }


def final_size_list(template_sizes, articles, cfg):
    """Größenspalten der Liste: Vorlage + benötigte Zusatzgrößen. Mit rules.remove_unused_sizes nur Größen,
    für die mindestens ein Artikel Bestand hat (weniger Spalten, übersichtlicher)."""
    used = {s for a in articles for s in a["stock"]}
    sizes = list(template_sizes)
    def number(s):
        try:
            return float(s)
        except ValueError:
            return None

    for extra in cfg["sizes"]["extra_sizes"]:
        if extra["size"] in used and extra["size"] not in sizes:
            if extra["insert_after"] in sizes:
                pos = sizes.index(extra["insert_after"]) + 1
            else:
                # Bezugsgröße fehlt (z.B. 37.5 ohne 37): numerisch einsortieren, sonst vor One-Size
                n = number(extra["size"])
                pos = next((i for i, s in enumerate(sizes) if n is not None and number(s) is not None and number(s) > n),
                           sizes.index(cfg["sizes"]["no_size_column"]) if cfg["sizes"]["no_size_column"] in sizes else len(sizes))
            sizes.insert(pos, extra["size"])
    if cfg["rules"].get("remove_unused_sizes"):
        in_stock = {s for a in articles for s, q in a["stock"].items() if q > 0}
        sizes = [s for s in sizes if s in in_stock] or sizes[:1]
    return sizes


def apply(cell, info, value=None, keep_template_value=False):
    cell._style = copy(info["style"])  # Kopie: spätere Änderungen an einer Zelle dürfen andere nicht mitändern
    if keep_template_value:
        cell.value = info["value"]
    elif value is not None:
        cell.value = value


def section_summary(articles):
    """Abschnitte in Listenreihenfolge: [(Titel, Anzahl Artikel), ...]."""
    out = OrderedDict()
    for art in articles:
        if art.get("section"):
            out[art["section"]] = out.get(art["section"], 0) + 1
    return list(out.items())


def nav_layout(sections, widths, first_col, last_col):
    """Verteilt die Sprungmarken auf Zeilen: je Link so viele Spalten, wie der Text braucht.
    Liefert [[(Titel, Text, Startspalte, Endspalte), ...], ...] (eine Liste je Zeile)."""
    rows, cur, col = [], [], first_col
    for title, n in sections:
        text = f"{title.upper()} ({n})"
        need, start, have = len(text) * 1.15 + 3, col, 0
        while have < need and col <= last_col:
            have += widths[col]
            col += 1
        if have < need and cur:  # passt nicht mehr in diese Zeile
            rows.append(cur)
            cur, col = [], first_col
            start, have = col, 0
            while have < need and col <= last_col:
                have += widths[col]
                col += 1
        cur.append((title, text, start, col - 1))
    if cur:
        rows.append(cur)
    return rows


DARK = "FF2F2F2F"
GREY_TEXT = "FF7F7F7F"
BAR_FILL = PatternFill("solid", fgColor="FFF5F5F5")
THIN = Side(style="thin", color="FFBFBFBF")


def write_top_bar(ws, top_row, n_rows, last_col):
    """Hellgraue Leiste über der Kopfzeile (Suche links, Inhalt rechts)."""
    for row in range(top_row, top_row + n_rows):
        ws.row_dimensions[row].height = 26
        for col in range(1, last_col + 1):
            ws.cell(row, col).fill = BAR_FILL


def write_nav_links(ws, top_row, nav_rows, section_rows):
    """Inhalt: Klick auf eine Gruppe springt zu ihrer Überschrift."""
    for i, links in enumerate(nav_rows):
        row = top_row + i
        for title, text, start, end in links:
            c = ws.cell(row, start, text)
            c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{ws.title}'!A{section_rows[title]}", display=text)
            c.font = Font(name="Arial", size=10, bold=True, underline="single", color=DARK)
            c.alignment = Alignment(horizontal="center", vertical="center")
            if end > start:
                ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)


def write_search(ws, row, scfg, fixed_cols, helper_col, first_row, last_row):
    """Suchfeld (SKU oder Namensteil) mit Sprung-Link; passende Artikel werden gelb markiert.
    Reine Formeln, keine Makros: funktioniert in Excel ohne Freigabe."""
    sku_l = get_column_letter(fixed_cols["sku"])
    name_l = get_column_letter(fixed_cols["name"])
    in_col, in_end = fixed_cols["sku"], fixed_cols["name"]
    res_col = min(fixed_cols["unit_price"], fixed_cols["uvp"])
    res_end = max(fixed_cols["unit_price"], fixed_cols["uvp"])
    q = f"${sku_l}${row}"
    rng_sku = f"${sku_l}${first_row}:${sku_l}${last_row}"
    rng_name = f"${name_l}${first_row}:${name_l}${last_row}"
    off = first_row - 1

    label = ws.cell(row, 1, scfg["label"])
    label.font = Font(name="Arial", size=10, bold=True, color=DARK)
    label.alignment = Alignment(horizontal="right", vertical="center", indent=1)

    box = ws.cell(row, in_col)
    box.font = Font(name="Arial", size=11, color=DARK)
    box.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    box.number_format = "@"
    for col in range(in_col, in_end + 1):
        ws.cell(row, col).fill = PatternFill("solid", fgColor="FFFFFFFF")
        ws.cell(row, col).border = Border(
            top=Side(style="medium", color=DARK), bottom=Side(style="medium", color=DARK),
            left=Side(style="medium", color=DARK) if col == in_col else Side(),
            right=Side(style="medium", color=DARK) if col == in_end else Side())
    ws.merge_cells(start_row=row, start_column=in_col, end_row=row, end_column=in_end)
    dv = DataValidation(type="textLength", operator="lessThan", formula1="100", allow_blank=True,
                        showInputMessage=True, promptTitle=scfg["prompt_title"][:32], prompt=scfg["prompt"][:255])
    dv.add(box.coordinate)
    ws.add_data_validation(dv)

    # Hilfszelle (ausgeblendete Spalte): Zeile des ersten Treffers – erst exakte SKU, dann SKU-Teil, dann Namensteil
    helper = ws.cell(row, helper_col)
    h_ref = f"${get_column_letter(helper_col)}${row}"
    helper.value = (
        f'=IF({q}="","",IFERROR(MATCH({q}&"",{rng_sku},0)+{off},'
        f'IFERROR(MATCH("*"&{q}&"*",{rng_sku},0)+{off},'
        f'IFERROR(MATCH("*"&{q}&"*",{rng_name},0)+{off},"-"))))'
    )
    ws.column_dimensions[get_column_letter(helper_col)].hidden = True

    res = ws.cell(row, res_col)
    res.value = (
        f'=IF({q}="","{scfg["empty_hint"]}",IF({h_ref}="-","{scfg["not_found"]}",'
        f'HYPERLINK("#\'{ws.title}\'!A"&{h_ref},"{scfg["jump"]} "&INDEX(${sku_l}:${sku_l},{h_ref}))))'
    )
    res.font = Font(name="Arial", size=10, bold=True, underline="single", color="FF1F4E79")
    res.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if res_end > res_col:
        ws.merge_cells(start_row=row, start_column=res_col, end_row=row, end_column=res_end)

    # Treffer gelb markieren (beide Zeilen eines Artikels)
    return FormulaRule(
        formula=[f'AND({q}<>"",OR(ISNUMBER(SEARCH({q},$' + sku_l + "{r}&\" \"&$" + name_l + "{r})),"
                 f"ISNUMBER(SEARCH({q},$" + sku_l + "{p}&\" \"&$" + name_l + "{p}))))"],
        fill=PatternFill("solid", fgColor="FFFFF2A8", bgColor="FFFFF2A8"),
    )


def write_legend(ws, row, first_col, last_col, widths, hints, stock_style, order_style):
    """Legende in Zeile 1: 'HOW TO ORDER' + zwei Farbchips in den Farben der Bestands- und Bestellzeile."""
    items = [(hints["legend_label"], None), (hints["legend_stock"], stock_style), (hints["legend_order"], order_style)]
    col = first_col
    for text, style in items:
        need, start, have = len(text) * 1.2 + 3, col, 0
        while have < need and col <= last_col:
            have += widths[col]
            col += 1
        c = ws.cell(row, start)
        if style is not None:
            for cc in range(start, col):
                apply(ws.cell(row, cc), style)
                ws.cell(row, cc).border = Border(top=THIN, bottom=THIN,
                                                 left=THIN if cc == start else Side(), right=THIN if cc == col - 1 else Side())
        c.value = text
        c.font = Font(name="Arial", size=9, bold=True, color=DARK if style is not None else GREY_TEXT)
        c.alignment = Alignment(horizontal="center" if style is not None else "right", vertical="center")
        if col - 1 > start:
            ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=col - 1)
        col += 1  # Abstand


def write_section_row(ws, row, title, last_col, n=None):
    """Überschrift einer Produktgruppe: über alle Spalten, fett, helles Grau, mit Artikelzahl."""
    cell = ws.cell(row, 1)
    parts = [TextBlock(InlineFont(rFont="Arial", sz=13, b=True, color=DARK), title.upper())]
    if n is not None:
        parts.append(TextBlock(InlineFont(rFont="Arial", sz=10, color=GREY_TEXT),
                               f"     {n} {'article' if n == 1 else 'articles'}"))
    cell.value = CellRichText(parts)
    cell.font = Font(name="Arial", size=13, bold=True, color=DARK)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    fill = PatternFill("solid", fgColor="FFEDEDED")
    for col in range(1, last_col + 1):
        ws.cell(row, col).fill = fill
        ws.cell(row, col).border = Border(top=Side(style="medium", color=DARK))
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.row_dimensions[row].height = 28


def write_workbook(template_path, out_path, articles, cfg, today, brand, show_discount=False):
    wb = load_workbook(template_path)
    ws = wb.worksheets[0]
    xc = cfg["excel"]
    layout = read_template_layout(ws, cfg)
    sizes = final_size_list(layout["sizes"], articles, cfg)
    roles = layout["roles"]
    hr, r_first = xc["header_row"], xc["first_data_row"]
    gcfg = cfg.get("product_groups", {})
    sections = section_summary(articles)

    # Reihenfolge der festen Spalten: Bild, SKU, Name, dann Preise wie in mapping.json (excel.price_order),
    # bei Discount zusätzlich DISCOUNT (%) und DISCOUNT PRICE (Stil wie UNIT PRICE)
    order = ["image", "sku", "name"] + list(xc.get("price_order", ["unit_price", "uvp"]))
    if show_discount:
        order += ["discount", "discount_price"]
        roles = dict(roles, discount=roles["unit_price"], discount_price=roles["unit_price"])
    layout = dict(layout, fixed_order=order, fixed_cols={r: i + 1 for i, r in enumerate(order)}, roles=roles)
    n_fixed = len(layout["fixed_order"])
    size_col = {s: n_fixed + 1 + i for i, s in enumerate(sizes)}
    q_col = n_fixed + len(sizes) + 1
    t_col = q_col + 1

    # Datenbereich ab Kopfzeile leeren (Titel in Zeile 1 bleibt erhalten)
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))
    ws.delete_rows(hr, ws.max_row)
    for key in list(ws.column_dimensions.keys()):
        del ws.column_dimensions[key]
    for r in [r for r in ws.row_dimensions if r >= hr]:
        del ws.row_dimensions[r]

    # Titel + Blattname
    date_str = today.strftime(xc["date_format"])
    if xc.get("title"):
        ws.cell(1, 1).value = xc["title"].format(BRAND=brand.upper(), brand=brand, date=date_str)
    elif layout["title"]:
        ws.cell(1, 1).value = str(layout["title"]).replace(xc["date_placeholder"], date_str)
    sheet_name = xc["sheet_name"].format(brand=brand, date=date_str)
    if len(sheet_name) > 31:
        # Excel erlaubt max. 31 Zeichen: Markenname kürzen
        cut = len(sheet_name) - 31
        sheet_name = xc["sheet_name"].format(brand=brand[:max(1, len(brand) - cut)].strip(), date=date_str)
    ws.title = sheet_name[:31]

    # Spaltenbreiten
    for role, col in layout["fixed_cols"].items():
        ws.column_dimensions[get_column_letter(col)].width = roles[role]["width"]
    for col in size_col.values():
        ws.column_dimensions[get_column_letter(col)].width = roles["size"]["width"]
    ws.column_dimensions[get_column_letter(q_col)].width = roles["quantity"]["width"]
    ws.column_dimensions[get_column_letter(t_col)].width = roles["total"]["width"]

    # Inhaltszeile(n) mit Sprungmarken zwischen Titel und Kopfzeile
    widths = {c: ws.column_dimensions[get_column_letter(c)].width for c in range(1, t_col + 1)}
    nav_rows = []
    if sections and gcfg.get("navigation", True):
        nav_rows = nav_layout(sections, widths, size_col[sizes[0]], t_col)
    scfg = xc.get("search") if xc.get("search", {}).get("enabled") else None
    bar_rows = max(len(nav_rows), 1 if scfg else 0)
    nav_top = hr
    hr, r_first = hr + bar_rows, r_first + bar_rows
    title_cell = ws.cell(1, 1)
    title_cell.font = Font(name="Arial", size=18, bold=True, color=DARK)
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 40

    # Kopfzeile
    ws.row_dimensions[hr].height = layout["row_heights"]["header"]
    for role, col in layout["fixed_cols"].items():
        apply(ws.cell(hr, col), roles[role]["header"], keep_template_value=True)
    for s, col in size_col.items():
        apply(ws.cell(hr, col), roles["size"]["header"], s)
    apply(ws.cell(hr, q_col), roles["quantity"]["header"], keep_template_value=True)
    apply(ws.cell(hr, t_col), roles["total"]["header"], keep_template_value=True)
    labels = dict({"discount": "DISCOUNT", "discount_price": "DISCOUNT PRICE"}, **xc.get("header_labels", {}))
    for role, col in list(layout["fixed_cols"].items()) + [("quantity", q_col), ("total", t_col)]:
        if labels.get(role):
            ws.cell(hr, col).value = labels[role]
    stock_labels = xc.get("stock_row_labels", {})

    first_size_letter = get_column_letter(size_col[sizes[0]])
    last_size_letter = get_column_letter(size_col[sizes[-1]])
    q_letter, t_letter = get_column_letter(q_col), get_column_letter(t_col)
    # Bestellsumme mit Discount-Preis, wenn vorhanden
    price_letter = get_column_letter(layout["fixed_cols"]["discount_price" if show_discount else "unit_price"])

    hints = xc.get("order_hints") if xc.get("order_hints", {}).get("enabled") else None
    first_article_row = None
    row = r_first
    section = None
    section_rows = {}
    counts = dict(sections)
    for art in articles:
        if art.get("section") and art["section"] != section:
            section = art["section"]
            write_section_row(ws, row, section, t_col, counts[section])
            section_rows[section] = row
            row += 1
        top, bottom = row, row + 1
        ws.row_dimensions[top].height = layout["row_heights"]["top"]
        ws.row_dimensions[bottom].height = layout["row_heights"]["bottom"]

        values = {"image": None, "sku": art["sku"], "name": art["name"],
                  "unit_price": float(art["unit_price"]), "uvp": float(art["uvp"]),
                  "discount": float(art["discount"]) if art.get("discount") else None,
                  "discount_price": float(art["discount_price"]) if "discount_price" in art else None}
        for role, col in layout["fixed_cols"].items():
            apply(ws.cell(top, col), roles[role]["top"], values[role])
            apply(ws.cell(bottom, col), roles[role]["bottom"])
        if show_discount:
            ws.cell(top, layout["fixed_cols"]["discount"]).number_format = "0%"
            dp = ws.cell(top, layout["fixed_cols"]["discount_price"])
            dp.font = Font(name=dp.font.name, size=dp.font.sz, bold=True, color=DARK)
        for s, col in size_col.items():
            qty = art["stock"].get(s)
            if qty == 0 and not cfg["rules"]["show_zero_stock"]:
                qty = None
            apply(ws.cell(top, col), roles["size"]["top"], qty)
            apply(ws.cell(bottom, col), roles["size"]["bottom"])
            if not art["stock"].get(s):
                # kein Bestand: keine orange Bestellzelle
                ws.cell(bottom, col).fill = PatternFill("solid", fgColor="FFFFFFFF")
        apply(ws.cell(top, q_col), roles["quantity"]["top"], keep_template_value=True)
        apply(ws.cell(top, t_col), roles["total"]["top"], keep_template_value=True)
        for role, col in (("quantity", q_col), ("total", t_col)):
            if stock_labels.get(role):
                ws.cell(top, col).value = stock_labels[role]
        apply(ws.cell(bottom, q_col), roles["quantity"]["bottom"])
        apply(ws.cell(bottom, t_col), roles["total"]["bottom"])
        for col in (q_col, t_col):
            # Ergebnisfelder (Formeln): nicht orange, damit Orange nur "hier eintragen" bedeutet
            ws.cell(bottom, col).fill = PatternFill("solid", fgColor="FFFFFFFF")
            ws.cell(bottom, col).font = Font(name="Arial", size=11, bold=True, color=DARK)
        if cfg["rules"]["order_formulas"]:
            ws.cell(bottom, q_col).value = (
                f'=IF(SUM({first_size_letter}{bottom}:{last_size_letter}{bottom})=0,"",'
                f"SUM({first_size_letter}{bottom}:{last_size_letter}{bottom}))"
            )
            ws.cell(bottom, t_col).value = f'=IF({q_letter}{bottom}="","",{q_letter}{bottom}*{price_letter}{top})'
        # Größen ohne Bestand: Eingabe gesperrt (max. 0); mit Bestand: Hinweis beim Klick, max. = Bestand darüber
        runs, run = [], None
        for s, col in size_col.items():
            has = bool(art["stock"].get(s))
            if run and run[0] == has and run[2] == col - 1:
                run[2] = col
            else:
                run = [has, col, col]
                runs.append(run)
        for has, c1, c2 in runs:
            l1, l2 = get_column_letter(c1), get_column_letter(c2)
            if has and hints:
                dv = DataValidation(
                    type="whole", operator="between", formula1="0",
                    formula2=f"{l1}{top}" if hints.get("limit_to_stock") else "100000",
                    allow_blank=True, showInputMessage=True, showErrorMessage=True,
                    promptTitle=hints["prompt_title"][:32], prompt=hints["prompt"][:255],
                    errorTitle=hints["error_title"][:32], error=hints["error"][:255],
                )
            elif not has and hints and hints.get("limit_to_stock"):
                dv = DataValidation(type="whole", operator="equal", formula1="0", allow_blank=True,
                                    showErrorMessage=True, errorTitle=hints["error_title"][:32],
                                    error=hints["error_none"][:255])
            else:
                continue
            dv.add(f"{l1}{bottom}:{l2}{bottom}" if c2 > c1 else f"{l1}{bottom}")
            ws.add_data_validation(dv)
        if first_article_row is None:
            first_article_row = top

        for col in layout["fixed_cols"].values():
            ws.merge_cells(start_row=top, start_column=col, end_row=bottom, end_column=col)
        if art.get("image"):
            # Platz = Spaltenbreite x Höhe beider Zeilen, abzüglich Rand
            box_w = int(roles["image"]["width"] * 8 + 5) - 8  # ca. 8 px je Zeichen bei Aptos Narrow 12
            box_h = int((layout["row_heights"]["top"] + layout["row_heights"]["bottom"]) * 4 / 3) - 8
            add_image(ws, art["image"], top, layout["fixed_cols"]["image"], box_w, box_h)
        row += 2

    last_row = row - 1
    if bar_rows:
        write_top_bar(ws, nav_top, bar_rows, t_col)
    if nav_rows:
        write_nav_links(ws, nav_top, nav_rows, section_rows)
    if scfg and first_article_row:
        rule = write_search(ws, nav_top, scfg, layout["fixed_cols"], t_col + 1, r_first, last_row)
        rng = f"A{r_first}:{t_letter}{last_row}"
        rule.formula = [rule.formula[0].replace("{r}", str(r_first)).replace("{p}", str(r_first - 1))]
        ws.conditional_formatting.add(rng, rule)
    if hints and first_article_row:
        # Legende direkt nach dem Titel (ab der Spalte hinter "Name"), damit sie auch bei wenigen Größen passt
        write_legend(ws, 1, layout["fixed_cols"]["name"] + 1, t_col, widths, hints,
                     roles["size"]["top"], roles["size"]["bottom"])
    if cfg["rules"]["grand_total_row"] and cfg["rules"]["order_formulas"]:
        ws.row_dimensions[row].height = layout["row_heights"]["bottom"]
        apply(ws.cell(row, q_col - 1), roles["quantity"]["bottom"], xc["grand_total_label"])
        apply(ws.cell(row, q_col), roles["quantity"]["bottom"],
              f"=SUM({q_letter}{r_first}:{q_letter}{last_row})")
        apply(ws.cell(row, t_col), roles["total"]["bottom"],
              f"=SUM({t_letter}{r_first}:{t_letter}{last_row})")
        for col in (q_col - 1, q_col, t_col):
            # Gesamtsumme ist ein Ergebnis, kein Eingabefeld: weiß und fett statt orange
            ws.cell(row, col).fill = PatternFill("solid", fgColor="FFFFFFFF")
            ws.cell(row, col).font = Font(name="Arial", size=11, bold=True, color=DARK)

    ws.freeze_panes = ws.cell(r_first, n_fixed + 1).coordinate
    if cfg["rules"]["landscape_print"]:
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape"
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.5
        ws.page_margins.header = ws.page_margins.footer = 0.3
        ws.print_options.horizontalCentered = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"1:{hr}"

    wb.save(out_path)
    return sizes


def output_path(cfg, today, brand):
    out_dir = BASE_DIR / cfg["paths"]["output_dir"]
    out_dir.mkdir(parents=True, exist_ok=True)
    safe_brand = re.sub(r"[^A-Za-z0-9]+", "-", normalize_name(brand).title()).strip("-")
    name = cfg["paths"]["output_filename"].format(brand=safe_brand, date_iso=today.isoformat())
    path = out_dir / name
    version = 2
    while path.exists():
        path = out_dir / f"{Path(name).stem}_v{version}{Path(name).suffix}"
        version += 1
    return path


def archive_csv(csv_path, cfg, today):
    archive_dir = BASE_DIR / cfg["paths"]["archive_dir"]
    archive_dir.mkdir(parents=True, exist_ok=True)
    target = archive_dir / f"{today.isoformat()}_{csv_path.name}"
    version = 2
    while target.exists():
        target = archive_dir / f"{today.isoformat()}_v{version}_{csv_path.name}"
        version += 1
    shutil.move(str(csv_path), target)
    return target


# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Excel-Stockliste aus CSV erzeugen")
    parser.add_argument("--csv", help="bestimmte CSV-Datei verwenden")
    parser.add_argument("--no-archive", action="store_true",
                        help="CSV nach Erfolg nicht nach input/archiv verschieben")
    parser.add_argument("--shop", help="Marke/Shop aus config/shopify.env (bei mehreren Shops)")
    parser.add_argument("--markup", help="EK-Preis = UVP ÷ Markup, wenn die CSV keinen EK-Preis hat (z.B. 2,5)")
    parser.add_argument("--sort", choices=[k for k, _ in SORT_CHOICES],
                        help="bestseller = Bestseller & Kategorien, categories = nur Kategorien, "
                             "stock = nach Bestand (sonst Abfrage)")
    args = parser.parse_args()

    cfg = load_config()
    today = date.today()
    warnings = []
    try:
        shop = select_shop(load_shops(cfg), args.shop)
        brand = shop["name"] if shop else cfg["excel"]["default_brand"]
        sort_label = select_sort(args.sort, cfg)
        apply_brand_groups(cfg, brand)
        print(f"Marke:    {brand}" + (f" ({shop['shop']})" if shop else ""))
        print(f"Sortierung: {sort_label}")
        csv_path = find_input_csv(cfg, args.csv)
        print(f"CSV:      {csv_path.name}")
        raw, enc = read_csv(csv_path, cfg)
        header_idx, positions, col_info = detect_columns(raw, cfg, brand, warnings)
        print(f"Encoding: {enc}, {len(raw) - header_idx - 1} Datenzeilen" + (f" (Kopfzeile in Zeile {header_idx + 1})" if header_idx else ""))
        print("Spalten:  " + "\n          ".join(col_info))
        markup = None
        if "unit_price" not in positions:
            markup = ask_markup(args.markup)
            print(f"EK-Preis: UVP ÷ {str(markup).replace('.', ',')} (Markup)")
        elif args.markup:
            warnings.append("--markup ignoriert: die CSV hat einen EK-Preis.")

        template_path = BASE_DIR / cfg["paths"]["template"]
        if not template_path.is_file():
            raise StocklistError(f"Template nicht gefunden: {template_path}")
        template_sizes = read_template_layout(load_workbook(template_path).worksheets[0], cfg)["sizes"]

        articles = build_articles(raw, positions, cfg, template_sizes, warnings, header_idx, brand, markup)
        enrich_from_shopify(articles, cfg, warnings, shop)
        sort_articles(articles, cfg, warnings)
        out = output_path(cfg, today, brand)
        sizes = write_workbook(template_path, out, articles, cfg, today, brand, "discount" in positions)
    except (StocklistError, OSError) as e:
        print("\nFEHLER – es wurde KEINE Stockliste erstellt.\n")
        print(e)
        for w in warnings:
            print(f"Hinweis: {w}")
        return 1

    for w in warnings:
        print(f"Hinweis: {w}")
    extra = [s for s in sizes if s not in template_sizes]
    if extra:
        print(f"Hinweis: zusätzliche Größenspalten eingefügt: {', '.join(extra)}")
    total = sum(sum(a["stock"].values()) for a in articles)
    print(f"\nFertig: {out.relative_to(BASE_DIR)}  ({len(articles)} Artikel, {total} Teile)")
    if not args.no_archive and not args.csv:
        moved = archive_csv(csv_path, cfg, today)
        print(f"CSV archiviert: {moved.relative_to(BASE_DIR)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
